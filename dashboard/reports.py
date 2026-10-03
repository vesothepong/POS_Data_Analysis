"""Report registry. Each report returns (columns, rows) so the same data feeds the page, CSV and XLSX export."""
import csv
import io
from decimal import Decimal

from django.http import HttpResponse

from analytics import calc
from forecasting.models import Forecast
from products.services import inventory_rows
from sales.models import SalesRecord
from sales.services import sales_dataframe

REPORT_TITLES = {
    "sales": "Sales Report",
    "inventory": "Inventory Report",
    "performance": "Product Performance Report",
    "reorder": "Reorder Report",
    "revenue": "Revenue Report",
}


def _sales():
    cols = ["Date", "Product", "Category", "Quantity Sold", "Price", "Total"]
    qs = SalesRecord.objects.select_related("product").order_by("-sale_date", "product__name")
    return cols, [[s.sale_date, s.product.name, s.product.category or "Uncategorized", s.quantity_sold, s.price, s.total_amount] for s in qs]


def _inventory():
    cols = ["Product", "Category", "Current Stock", "Reorder Level", "Predicted Demand", "Reorder Quantity", "Status"]
    return cols, [[r["product"].name, r["product"].category or "Uncategorized", r["stock"], r["reorder_level"],
                   r["predicted"] if r["predicted"] is not None else "", r["reorder_qty"], r["status"]] for r in inventory_rows()]


def _performance():
    cols = ["Rank", "Product", "Category", "Units Sold", "Revenue", "Avg Units / Month"]
    df = sales_dataframe()
    if df.empty:
        return cols, []
    months = df.groupby("product")["month"].nunique()
    cat = df.groupby("product")["category"].first()
    rows = []
    for i, r in enumerate(calc.group_totals(df, "product"), start=1):
        rows.append([i, r["name"], cat[r["name"]], r["units"], r["revenue"], round(r["units"] / max(1, months[r["name"]]), 1)])
    return cols, rows


def _reorder():
    cols = ["Product", "Current Stock", "Predicted Demand", "Reorder Quantity", "Status"]
    return cols, [[r["product"].name, r["stock"], r["predicted"], r["reorder_qty"], r["status"]]
                  for r in inventory_rows() if r["reorder_qty"] > 0]


def _revenue():
    cols = ["Month", "Units Sold", "Revenue"]
    monthly_data = calc.monthly_totals(sales_dataframe(), ascending=False)
    rows = [[m["month"], m["units"], m["revenue"]] for m in monthly_data]
    if rows:
        total_units = sum(m["units"] for m in monthly_data)
        total_rev = round(sum(m["revenue"] for m in monthly_data), 2)
        rows.append(["TOTAL", total_units, total_rev])
    return cols, rows


BUILDERS = {"sales": _sales, "inventory": _inventory,
            "performance": _performance, "reorder": _reorder, "revenue": _revenue}


def build(kind):
    return BUILDERS[kind]()


def _plain(v):
    if isinstance(v, Decimal):
        return float(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10] if hasattr(v, "year") and not hasattr(v, "hour") else str(v)
    return v


def export(kind, fmt):
    cols, rows = build(kind)
    name = f"{kind}_report"
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
        w = csv.writer(resp)
        w.writerow(cols)
        for r in rows:
            w.writerow([_plain(v) for v in r])
        return resp
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = REPORT_TITLES[kind][:31]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_plain(v) for v in r])
    for i, col in enumerate(cols, start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 40)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp
